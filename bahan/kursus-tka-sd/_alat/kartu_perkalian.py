"""Membangun kartu fakta perkalian 1–10 (A4, huruf besar) -> _bangun/kartu/kartu-fakta-perkalian.docx"""
import os
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import md_docx as M

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = os.path.join(ROOT, "_bangun", "kartu")
os.makedirs(out, exist_ok=True)

d = M.dokumen_baru(ukuran_dasar=14, margin_cm=1.5)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Kartu Fakta Perkalian 1–10")
r.bold, r.font.size = True, Pt(26)
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Cari baris bilangan pertama, lalu kolom bilangan kedua. Contoh: 7 × 8 → baris 7, kolom 8 = 56")
r.font.size = Pt(13)

t = d.add_table(rows=11, cols=11)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i in range(11):
    t.rows[i].height = Cm(2.0)
    for j in range(11):
        c = t.cell(i, j)
        c.width = Cm(1.65)
        if i == 0 and j == 0:
            teks = "×"
        elif i == 0:
            teks = str(j)
        elif j == 0:
            teks = str(i)
        else:
            teks = str(i * j)
        para = c.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(teks)
        run.font.size = Pt(20 if (i == 0 or j == 0) else 18)
        run.bold = (i == 0 or j == 0 or i == j)
        if i == 0 or j == 0:
            M.arsir(c._tc.get_or_add_tcPr(), "D9D9D9")

p = d.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Untuk pengajar: kartu ini untuk anak dengan skor tes awal Matematika 7–11, minggu 1–2. "
              "Boleh dilihat saat latihan dan tiket keluar. Tidak dipakai saat kuis soal fakta perkalian (PR-M2). "
              "Letakkan di meja anak sebelum kelas, tanpa diumumkan.")
r.font.size = Pt(11)
d.save(os.path.join(out, "kartu-fakta-perkalian.docx"))
print("ok")
