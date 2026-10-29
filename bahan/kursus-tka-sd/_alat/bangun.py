"""Tahap 5 — membangun berkas siap pakai.

python bangun.py modul        -> 06-siap-pakai/modul-docx/*.docx
python bangun.py lembar       -> _bangun/lembar/*.docx  (lalu dikonversi ke lembar-kerja-anak/*.pdf)
python bangun.py simulasi     -> _bangun/simulasi/*.docx (lalu ke simulasi/*.pdf)
python bangun.py buku         -> _bangun/buku/buku-pegangan-pengajar.docx
python bangun.py semua
Konversi PDF: powershell -File konversi.ps1 (Word COM).
"""
import os, re, sys, glob
from docx.shared import Pt, Cm
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import md_docx as M
import periksa as P

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SIAP = os.path.join(ROOT, "06-siap-pakai")
BANGUN = os.path.join(ROOT, "_bangun")
SIM_KODE = {"M08", "M14", "M20", "M24", "B12"}


def p_of(kode):
    n = int(kode[1:])
    if kode[0] == "B":
        return 3 * n
    k = (n + 1) // 2
    return 3 * k - 2 if n % 2 else 3 * k - 1


def modul_semua():
    return sorted(glob.glob(os.path.join(ROOT, "01-modul-matematika", "M??-*.md"))) + \
        sorted(glob.glob(os.path.join(ROOT, "02-modul-bahasa-indonesia", "B??-*.md")))


def teks_T(P_):
    f = glob.glob(os.path.join(ROOT, "03-bank-teks-baca-cepat", f"T{P_:02d}-*.md"))
    return (f[0], open(f[0], encoding="utf-8").read()) if f else (None, None)


def judul_md(teks):
    m = re.search(r"^# (.+)$", teks, re.M)
    return m.group(1).strip() if m else ""


def bagian(teks, nama):
    """Isi di bawah '## nama' sampai '## ' berikutnya."""
    m = re.search(rf"^## {re.escape(nama)}\s*$", teks, re.M)
    if not m:
        return ""
    sisa = teks[m.end():]
    n = re.search(r"^## ", sisa, re.M)
    return sisa[:n.start()] if n else sisa


def buang_kunci(teks):
    return "\n".join(l for l in teks.splitlines() if not l.strip().startswith(">"))


# ---------------------------------------------------------------- modul pengajar
def bangun_modul():
    out = os.path.join(SIAP, "modul-docx")
    os.makedirs(out, exist_ok=True)
    for f in modul_semua():
        teks = open(f, encoding="utf-8").read()
        kode = os.path.basename(f)[:3]
        d = M.dokumen_baru(ukuran_dasar=12, margin_cm=1.8)
        judul = judul_md(teks)
        M.nomor_halaman(d, judul)
        M.tulis_markdown(d, teks, ukuran=12, pecah_halaman_pada=r"Halaman \d")
        tf, tt = teks_T(p_of(kode))
        if tt:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = d.add_heading(level=2)
            M.tambah_inline(h, "Lampiran — Teks baca cepat " + judul_md(tt))
            M.tulis_markdown(d, tt, ukuran=12, lewati_judul1=True, geser_judul=1)
        d.save(os.path.join(out, os.path.basename(f)[:-3] + ".docx"))
        print("modul", kode)


# ---------------------------------------------------------------- lembar kerja anak
def soal_badan_ringkas(s, nomor_baru, peta_nomor):
    badan = s["badan"]
    badan = re.sub(r"(soal nomor|Soal nomor|soal no\.)\s*(\d+)",
                   lambda m: f"{m.group(1)} {peta_nomor.get(int(m.group(2)), m.group(2))}", badan)
    return badan


def isi_lembar(d, judul_bagian, soal_list, ukuran, ruang_kerja, nomor_awal=1, peta_nomor=None, pakai_nomor_modul=False,
               pengantar=None):
    h = d.add_heading(level=2)
    M.tambah_inline(h, judul_bagian)
    if pengantar:
        r = d.add_paragraph().add_run(pengantar)
        r.italic, r.font.size = True, Pt(12)
    for k, s in enumerate(soal_list):
        # nomor lembar anak = nomor di modul, supaya kunci yang dibacakan pengajar cocok
        n = s["nomor"] if pakai_nomor_modul else nomor_awal + k
        p = d.add_paragraph()
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"Soal {n}")
        r.bold = True
        r.font.size = Pt(ukuran + 1)
        perintah = {"PG kompleks": "  (pilih semua jawaban yang benar)",
                    "PG kategori": "  (tentukan Benar atau Salah)"}.get(s["bentuk"], "")
        if perintah:
            r2 = p.add_run(perintah)
            r2.font.size = Pt(ukuran - 2)
        M.tulis_markdown(d, soal_badan_ringkas(s, n, peta_nomor or {}), ukuran=ukuran)
        for _ in range(ruang_kerja):
            q = d.add_paragraph()
            q.paragraph_format.space_after = Pt(6)
        garis = d.add_paragraph()
        pPr = garis._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k2, v in (("w:val", "dotted"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
            b.set(qn(k2), v)
        bdr.append(b)
        pPr.insert_element_before(bdr, "w:shd", *M.PPR_SETELAH_SHD)


def kepala_lembar(d, judul, sub):
    t = d.add_paragraph()
    r = t.add_run(judul)
    r.bold = True
    r.font.size = Pt(18)
    s = d.add_paragraph()
    r = s.add_run(sub)
    r.font.size = Pt(12)
    n = d.add_paragraph()
    n.paragraph_format.space_before = Pt(6)
    r = n.add_run("Nama: ______________________________      Tanggal: ________________")
    r.font.size = Pt(14)


def bangun_lembar():
    out = os.path.join(BANGUN, "lembar")
    os.makedirs(out, exist_ok=True)
    for f in modul_semua():
        teks = open(f, encoding="utf-8").read()
        kode = os.path.basename(f)[:3]
        Pn = p_of(kode)
        mat = kode.startswith("M")
        judul = judul_md(teks)
        d = M.dokumen_baru(ukuran_dasar=14, margin_cm=1.8)
        M.nomor_halaman(d, f"Lembar kerja {kode} · P{Pn:02d}")
        kepala_lembar(d, f"Lembar Kerja — Pertemuan {Pn}", judul.split("—", 1)[-1].strip())
        # A. baca cepat
        tf, tt = teks_T(Pn)
        if tt:
            h = d.add_heading(level=2)
            M.tambah_inline(h, "A. Baca cepat — " + judul_md(tt).split("—", 1)[-1].strip())
            p = d.add_paragraph()
            r = p.add_run("Baca pertanyaannya dulu, baru teksnya.      Waktu selesai: ______ menit ______ detik")
            r.italic = True
            r.font.size = Pt(12)
            M.tulis_markdown(d, bagian(tt, "Teks"), ukuran=14)
            pert = [s for s in P.blok_soal(tt) if s["jenis"] == "Pertanyaan"]
            isi_lembar(d, "Pertanyaan baca cepat", pert, 14, 0)
        # B. latihan
        hal = P.potong_halaman(teks)
        if kode in ("M01", "B01"):
            h2 = hal.get(P.HALAMAN[1], "")
            diag = [s for s in P.blok_soal(h2) if s["jenis"] == "Diagnostik"]
            if diag:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                isi_lembar(d, "B. Tes awal", diag, 14, 1 if mat else 0)
        elif kode not in SIM_KODE:
            h5 = hal.get(P.HALAMAN[4], "")
            semua5 = [s for s in P.blok_soal(h5) if s["jenis"] == "Soal"]
            isi = [s for s in semua5 if s["lambang"] == "●"]
            if isi:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                isi_lembar(d, "B. Latihan campur-aduk", isi, 14, 3 if mat else 0, pakai_nomor_modul=True)
        # C. tiket keluar
        h6 = hal.get(P.HALAMAN[5], "")
        tiket = [s for s in P.blok_soal(h6) if s["jenis"] == "Tiket"]
        if tiket:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = d.add_paragraph("✂ - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - -")
            p = d.add_paragraph()
            r = p.add_run(f"Tiket keluar — Pertemuan {Pn}      Nama: ______________________")
            r.bold = True
            r.font.size = Pt(14)
            isi_lembar(d, "C. Tiket keluar (kerjakan sendiri)", tiket, 14, 2 if mat else 0)
        # D. latihan rumah: soal ○ + kunci (bukan M01/B01, bukan simulasi)
        if kode not in ("M01", "B01") and kode not in SIM_KODE:
            h5 = hal.get(P.HALAMAN[4], "")
            cadangan = [s for s in P.blok_soal(h5) if s["jenis"] == "Soal" and s["lambang"] == "○"]
            if cadangan:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                isi_lembar(d, "D. Latihan rumah", cadangan, 14, 3 if mat else 0, pakai_nomor_modul=True,
                           pengantar="Kerjakan di rumah ±10 menit: soal di bagian B yang belum selesai, lalu soal di bawah ini. "
                                     "Cocokkan sendiri dengan kunci di halaman terakhir. Tidak dinilai.")
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                h = d.add_heading(level=2)
                M.tambah_inline(h, "Kunci latihan rumah")
                p = d.add_paragraph()
                r = p.add_run("Kalau jawabanmu belum sama, baca Petunjuk dulu dan coba lagi sebelum melihat hurufnya.")
                r.italic, r.font.size = True, Pt(12)
                baris = ["| Soal | Petunjuk | Kunci |", "|---|---|---|"]
                for s in cadangan:
                    pet = s["kunci"].get("Petunjuk", "").replace("|", "/")
                    baris.append(f"| {s['nomor']} | {pet} | {s['kunci'].get('Kunci', '')} |")
                M.tulis_markdown(d, "\n".join(baris), ukuran=13)
                p = d.add_paragraph()
                r = p.add_run("Untuk orang tua: cukup ingatkan waktunya. Tidak perlu mengajari atau memeriksa jawaban.")
                r.font.size = Pt(11)
        d.save(os.path.join(out, f"LK-P{Pn:02d}-{kode}.docx"))
        print("lembar", kode)


# ---------------------------------------------------------------- simulasi
def bangun_simulasi():
    out = os.path.join(BANGUN, "simulasi")
    os.makedirs(out, exist_ok=True)
    fs = sorted(glob.glob(os.path.join(ROOT, "01-modul-matematika", "simulasi", "S*.md")) +
                glob.glob(os.path.join(ROOT, "02-modul-bahasa-indonesia", "simulasi", "S*.md")))
    for f in fs:
        teks = open(f, encoding="utf-8").read()
        nama = os.path.basename(f)[:-3]
        judul = judul_md(teks)
        # naskah anak: petunjuk + soal tanpa kunci + lembar jawaban
        d = M.dokumen_baru(ukuran_dasar=13, margin_cm=1.8)
        M.nomor_halaman(d, judul)
        kepala_lembar(d, judul, "30 soal · 75 menit")
        for sec, pecah in (("Petunjuk ujian", False), ("Naskah soal", True), ("Lembar jawaban", True)):
            isi = bagian(teks, sec)
            if not isi:
                continue
            if pecah:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = d.add_heading(level=2)
            M.tambah_inline(h, sec)
            isi = buang_kunci(isi)
            isi = re.sub(r"^### Stimulus\s+([A-Z])\s*:", r"### Teks \1:", isi, flags=re.M)
            isi = re.sub(r"\b[Ss]timulus\s+([A-Z])\b", r"teks \1", isi)
            isi = re.sub(r"^#### Butir (\d+)\s*·\s*(.+)$", lambda m: f"#### Soal {m.group(1)}" +
                         {"PG kompleks": " (pilih semua jawaban yang benar)",
                          "PG kategori": " (tentukan Benar atau Salah)"}.get(m.group(2).strip(), ""), isi, flags=re.M)
            M.tulis_markdown(d, isi, ukuran=13, ukuran_kunci=11)
        d.save(os.path.join(out, nama + "-naskah.docx"))
        # kunci pengajar: dokumen lengkap
        k = M.dokumen_baru(ukuran_dasar=11, margin_cm=1.8)
        M.nomor_halaman(k, judul + " — KUNCI (untuk pengajar)")
        M.tulis_markdown(k, teks.replace("# " + judul, "# " + judul + " — KUNCI PENGAJAR", 1), ukuran=11,
                         pecah_halaman_pada=r"(Naskah soal|Kunci ringkas|Pemetaan)")
        k.save(os.path.join(out, nama + "-kunci.docx"))
        print("simulasi", nama)


# ---------------------------------------------------------------- buku pegangan
def tambah_daftar_isi(d):
    p = d.add_paragraph()
    r = p.add_run()
    for tag, isi in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = isi
        r._r.append(el)


def bangun_buku():
    out = os.path.join(BANGUN, "buku")
    os.makedirs(out, exist_ok=True)
    fs = sorted(glob.glob(os.path.join(ROOT, "04-petunjuk", "0*.md")))
    d = M.dokumen_baru(ukuran_dasar=12, margin_cm=2.2)
    M.nomor_halaman(d, "Buku Pegangan Pengajar · Paket Siap TKA SD")
    for _ in range(8):
        d.add_paragraph()
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Buku Pegangan Pengajar")
    r.bold = True
    r.font.size = Pt(30)
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Paket Siap TKA SD · Kelas 6 · 36 pertemuan")
    r.font.size = Pt(16)
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Banda Aceh · 2026")
    r.font.size = Pt(12)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = d.add_paragraph()
    r = h.add_run("Daftar Isi")
    r.bold = True
    r.font.size = Pt(20)
    tambah_daftar_isi(d)
    for f in fs:
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        M.tulis_markdown(d, open(f, encoding="utf-8").read(), ukuran=12)
    d.save(os.path.join(out, "buku-pegangan-pengajar.docx"))
    print("buku", len(fs), "berkas")


def bangun_word():
    """Berkas Word terpisah per jenis: 10-berkas-word/{01-modul-matematika,02-modul-bahasa-indonesia,03-bank-teks-baca-cepat}."""
    dasar = os.path.join(ROOT, "10-berkas-word")
    tujuan = {"M": os.path.join(dasar, "01-modul-matematika"), "B": os.path.join(dasar, "02-modul-bahasa-indonesia")}
    for t in tujuan.values():
        os.makedirs(os.path.join(t, "paket-simulasi"), exist_ok=True)
    # modul
    for f in modul_semua():
        teks = open(f, encoding="utf-8").read()
        kode = os.path.basename(f)[:3]
        d = M.dokumen_baru(ukuran_dasar=12, margin_cm=1.8)
        M.nomor_halaman(d, judul_md(teks))
        M.tulis_markdown(d, teks, ukuran=12, pecah_halaman_pada=r"Halaman \d")
        d.save(os.path.join(tujuan[kode[0]], os.path.basename(f)[:-3] + ".docx"))
    # paket simulasi lengkap dengan kunci (untuk pengajar)
    for folder, huruf in (("01-modul-matematika", "M"), ("02-modul-bahasa-indonesia", "B")):
        for f in sorted(glob.glob(os.path.join(ROOT, folder, "simulasi", "S*.md"))):
            teks = open(f, encoding="utf-8").read()
            d = M.dokumen_baru(ukuran_dasar=12, margin_cm=1.8)
            M.nomor_halaman(d, judul_md(teks))
            M.tulis_markdown(d, teks, ukuran=12, pecah_halaman_pada=r"(Naskah soal|Lembar jawaban|Kunci ringkas)")
            d.save(os.path.join(tujuan[huruf], "paket-simulasi", os.path.basename(f)[:-3] + ".docx"))
    # bank teks: satu berkas per teks + satu berkas gabungan
    bt = os.path.join(dasar, "03-bank-teks-baca-cepat")
    os.makedirs(bt, exist_ok=True)
    fs = sorted(glob.glob(os.path.join(ROOT, "03-bank-teks-baca-cepat", "T??-*.md")))
    gab = M.dokumen_baru(ukuran_dasar=13, margin_cm=1.8)
    M.nomor_halaman(gab, "Bank Teks Baca Cepat T01–T36")
    for i, f in enumerate(fs):
        teks = open(f, encoding="utf-8").read()
        d = M.dokumen_baru(ukuran_dasar=13, margin_cm=1.8)
        M.nomor_halaman(d, judul_md(teks))
        M.tulis_markdown(d, teks, ukuran=13)
        d.save(os.path.join(bt, os.path.basename(f)[:-3] + ".docx"))
        if i:
            gab.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        M.tulis_markdown(gab, teks, ukuran=13)
    gab.save(os.path.join(bt, "00-bank-teks-baca-cepat-lengkap.docx"))
    # lembar kerja anak
    import shutil
    bangun_lembar()
    lk = os.path.join(dasar, "04-lembar-kerja-anak")
    os.makedirs(lk, exist_ok=True)
    for f in sorted(glob.glob(os.path.join(BANGUN, "lembar", "LK-*.docx"))):
        shutil.copy2(f, lk)
    # buku pegangan pengajar: satu buku + 7 petunjuk terpisah
    bangun_buku()
    bp = os.path.join(dasar, "05-buku-pegangan-pengajar")
    os.makedirs(os.path.join(bp, "petunjuk-terpisah"), exist_ok=True)
    shutil.copy2(os.path.join(BANGUN, "buku", "buku-pegangan-pengajar.docx"), bp)
    for f in sorted(glob.glob(os.path.join(ROOT, "04-petunjuk", "0*.md"))):
        teks = open(f, encoding="utf-8").read()
        d = M.dokumen_baru(ukuran_dasar=12, margin_cm=2.0)
        M.nomor_halaman(d, judul_md(teks))
        M.tulis_markdown(d, teks, ukuran=12)
        d.save(os.path.join(bp, "petunjuk-terpisah", os.path.basename(f)[:-3] + ".docx"))
    print("buku pegangan: jalankan perbarui_daftar_isi.py pada", os.path.join(bp, "buku-pegangan-pengajar.docx"))
    # bahan pelatihan pengajar: buku fasilitator + kartu peserta + berkas terpisah
    bangun_pelatihan()
    pl = os.path.join(dasar, "06-bahan-pelatihan-pengajar")
    os.makedirs(os.path.join(pl, "berkas-terpisah"), exist_ok=True)
    for nama in ("bahan-pelatihan-pengajar.docx", "kartu-dan-lembar-peserta.docx"):
        shutil.copy2(os.path.join(BANGUN, "pelatihan", nama), pl)
    for f in sorted(glob.glob(os.path.join(ROOT, "07-pelatihan-pengajar", "0[0-5]*.md"))):
        teks = open(f, encoding="utf-8").read()
        kartu = os.path.basename(f).startswith("06-")
        d = M.dokumen_baru(ukuran_dasar=13 if kartu else 12, margin_cm=1.8 if kartu else 2.0)
        M.nomor_halaman(d, judul_md(teks))
        M.tulis_markdown(d, teks, ukuran=13 if kartu else 11.5, pecah_halaman_pada=r".+" if kartu else None)
        d.save(os.path.join(pl, "berkas-terpisah", os.path.basename(f)[:-3] + ".docx"))
    print("bahan pelatihan: jalankan perbarui_daftar_isi.py pada", os.path.join(pl, "bahan-pelatihan-pengajar.docx"))
    # materi promosi
    pr = os.path.join(dasar, "07-materi-promosi")
    os.makedirs(pr, exist_ok=True)
    for f in sorted(glob.glob(os.path.join(ROOT, "09-promosi", "0*.md"))):
        teks = open(f, encoding="utf-8").read()
        d = M.dokumen_baru(ukuran_dasar=12, margin_cm=2.0)
        M.nomor_halaman(d, judul_md(teks))
        M.tulis_markdown(d, teks, ukuran=12)
        d.save(os.path.join(pr, os.path.basename(f)[:-3] + ".docx"))
    print("word:", len(modul_semua()), "modul,", len(fs), "teks")


def bangun_pelatihan():
    out = os.path.join(BANGUN, "pelatihan")
    os.makedirs(out, exist_ok=True)
    folder = os.path.join(ROOT, "07-pelatihan-pengajar")
    # buku fasilitator: 00–05
    d = M.dokumen_baru(ukuran_dasar=12, margin_cm=2.0)
    M.nomor_halaman(d, "Bahan Pelatihan Pengajar · Paket Siap TKA SD")
    for _ in range(8):
        d.add_paragraph()
    for teks, sz, tebal in (("Bahan Pelatihan Pengajar", 30, True), ("12 jam · 4 sesi · naskah fasilitator dan ujian kelulusan", 15, False),
                            ("Paket Siap TKA SD · Banda Aceh · 2026", 12, False)):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(teks)
        r.bold, r.font.size = tebal, Pt(sz)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    r = d.add_paragraph().add_run("Daftar Isi")
    r.bold, r.font.size = True, Pt(20)
    tambah_daftar_isi(d)
    for f in sorted(glob.glob(os.path.join(folder, "0[0-5]-*.md"))):
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        M.tulis_markdown(d, open(f, encoding="utf-8").read(), ukuran=11.5)
    d.save(os.path.join(out, "bahan-pelatihan-pengajar.docx"))
    # kartu dan lembar peserta: 06, tiap '## ' di halaman baru
    k = glob.glob(os.path.join(folder, "06-*.md"))
    if k:
        d = M.dokumen_baru(ukuran_dasar=13, margin_cm=1.8)
        M.nomor_halaman(d, "Kartu dan Lembar Peserta · Pelatihan Pengajar")
        M.tulis_markdown(d, open(k[0], encoding="utf-8").read(), ukuran=13, pecah_halaman_pada=r".+")
        d.save(os.path.join(out, "kartu-dan-lembar-peserta.docx"))
    print("pelatihan selesai")


if __name__ == "__main__":
    apa = sys.argv[1] if len(sys.argv) > 1 else "semua"
    if apa == "pelatihan":
        bangun_pelatihan()
    if apa == "word":
        bangun_word()
    if apa in ("modul", "semua"):
        bangun_modul()
    if apa in ("lembar", "semua"):
        bangun_lembar()
    if apa in ("simulasi", "semua"):
        bangun_simulasi()
    if apa in ("buku", "semua"):
        bangun_buku()
