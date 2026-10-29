"""Membangun 06-siap-pakai/rapor-penguasaan.xlsx (Tahap 5)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "06-siap-pakai", "rapor-penguasaan.xlsx")
N = 8              # baris murid
R0 = 5             # baris murid pertama
F = "Arial"
HIJAU, KUNING, MERAH = "C6EFCE", "FFEB9C", "FFC7CE"
ABU = "D9D9D9"
INPUT = "FFFFCC"
tipis = Side(style="thin", color="BFBFBF")
KOTAK = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)

# P -> (kode, topik singkat)
JADWAL = {}
MAT = ["Diagnostik & operasi hitung campuran", "Faktor, kelipatan, KPK, FPB", "Pecahan senilai & membandingkan",
       "Pecahan, desimal, persen", "Tambah & kurang pecahan", "Kali & bagi pecahan", "Operasi campuran lintas bentuk",
       "SIMULASI 1", "Satuan panjang, berat, volume", "Waktu & kecepatan", "Keliling & luas", "Sudut & kekongruenan",
       "Penaksiran & pembulatan", "SIMULASI 2", "Volume kubus & balok", "Luas permukaan & bangun gabungan",
       "Visualisasi spasial", "Rasio & perbandingan", "Senilai, berbalik nilai, skala", "SIMULASI 3",
       "Piktogram, diagram batang, tabel", "Menafsirkan data & rata-rata", "Latihan PG kompleks & kategori",
       "Latihan bertempo + simulasi akhir"]
BI = ["Diagnostik & baca pertanyaan dulu", "Informasi tersurat & kosakata", "Menyusun ulang informasi",
      "Ide pokok", "Tokoh, peristiwa, amanat", "Perubahan tokoh & latar", "Ungkapan & majas", "Fakta & opini",
      "Infografik, grafik, tabel", "Relevansi & kesesuaian", "Puisi & perasaan", "Simulasi akhir BI"]
for j, t in enumerate(MAT, 1):
    k = (j + 1) // 2
    P = 3 * k - 2 if j % 2 else 3 * k - 1
    JADWAL[P] = (f"M{j:02d}", t)
for k, t in enumerate(BI, 1):
    JADWAL[3 * k] = (f"B{k:02d}", t)
P_OF = {v[0]: p for p, v in JADWAL.items()}

# topik isi -> sumber uji tunda (dari 02-peta-spiral.md)
UJI = {"M01": "P13 · M09", "M02": "P14 · M10", "M03": "P16 · M11", "M04": "P17 · M12", "M05": "P19 · M13",
       "M06": "P22 · M15", "M07": "P23 · M16", "M09": "P25 · M17", "M10": "P26 · M18", "M11": "P28 · M19",
       "M12": "P31 · M21", "M13": "P32 · M22", "M15": "P34 · M23", "M16": "P35 · simulasi S4", "M17": "P35 · simulasi S4",
       "M18": "P35 · simulasi S4", "M19": "P35 · simulasi S4", "M21": "P35 · simulasi S4", "M22": "P35 · simulasi S4",
       "B01": "P15 · B05", "B02": "P18 · B06", "B03": "P21 · B07", "B04": "P24 · B08", "B05": "P27 · B09",
       "B06": "P30 · B10", "B07": "P33 · B11", "B08": "P36 · simulasi S5", "B09": "P36 · simulasi S5",
       "B10": "P36 · simulasi S5", "B11": "P36 · simulasi S5"}
TOPIK = list(UJI.keys())


def gaya(ws, lebar=None):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size if c.font and c.font.size else 10,
                              bold=c.font.bold if c.font else False, italic=c.font.italic if c.font else False,
                              color=c.font.color if c.font else None)


def judul(ws, teks, sub):
    ws["A1"] = teks
    ws["A1"].font = Font(name=F, size=14, bold=True)
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=9, italic=True)


def kepala(c, isi, putar=False):
    c.value = isi
    c.font = Font(name=F, size=9, bold=True)
    c.fill = PatternFill("solid", fgColor=ABU)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center",
                            text_rotation=90 if putar else 0)
    c.border = KOTAK


def nama_murid(ws, col=1):
    for i in range(N):
        r = R0 + i
        c = ws.cell(r, col, f"=IF('Data Murid'!B{r}=\"\",\"\",'Data Murid'!B{r})")
        c.border = KOTAK


wb = Workbook()

# ---------------------------------------------------------------- Petunjuk
ws = wb.active
ws.title = "Petunjuk"
baris = [
    ("Rapor Penguasaan — Paket Siap TKA SD", True, 14),
    ("Satu berkas untuk satu kelas (maks 8 murid). Salin berkas ini untuk kelas berikutnya.", False, 10),
    ("", False, 10),
    ("SEL YANG DIISI PENGAJAR — berlatar kuning muda:", True, 11),
    ("• Data Murid: nama, kelas asal, sekolah, kontak orang tua, skor diagnostik (baris 5 berisi CONTOH — ganti dengan murid sungguhan).", False, 10),
    ("• Tiket Keluar: skor 0–3 setiap pertemuan. Warna muncul sendiri: 3 hijau · 2 kuning · 0–1 merah.", False, 10),
    ("• Penguasaan Topik: hanya kolom 'Uji tunda' (0–3). Kolom 'Tiket' dan 'Status' terisi sendiri.", False, 10),
    ("• Kehadiran: H = hadir, T = tidak hadir. Persentase terhitung sendiri.", False, 10),
    ("• Survei Murid: pertanyaan 1–3 diisi angka 1–5, pertanyaan 4 diisi kalimat anak.", False, 10),
    ("• Laporan Orang Tua: tidak diisi — tinggal disalin ke WhatsApp.", False, 10),
    ("", False, 10),
    ("ATURAN WARNA PENGUASAAN (01-spesifikasi-program.md Bagian 8):", True, 11),
    ("Kuasai (hijau) = uji tunda 3 dari 3 — anak masih bisa mengerjakan topik itu ±3 minggu kemudian.", False, 10),
    ("Baru bisa (kuning) = tiket keluar ≥ 2 tetapi uji tunda belum ada, atau uji tunda = 2.", False, 10),
    ("Belum (merah) = tiket keluar ≤ 1 dan belum ada uji tunda, atau uji tunda ≤ 1.", False, 10),
    ("Uji tunda diambil dari 1 soal kuis 'bulan lalu' + 2 soal ● 'bulan lalu' di pertemuan yang tertulis di kepala kolom.", False, 10),
    ("Untuk M16–M22 dan B08–B11, uji tunda diambil dari butir simulasi akhir (aturan skala di 02-peta-spiral.md Bagian 6).", False, 10),
    ("", False, 10),
    ("YANG SENGAJA TIDAK ADA: peringkat, nilai rata-rata kelas, perbandingan antar-anak. Jangan ditambahkan.", True, 10),
    ("Dua kali merah berturut-turut di Tiket Keluar → telepon orang tua malam itu.", True, 10),
]
for i, (t, b, s) in enumerate(baris, 1):
    ws.cell(i, 1, t).font = Font(name=F, bold=b, size=s)
ws.column_dimensions["A"].width = 120

# ---------------------------------------------------------------- Data Murid
ws = wb.create_sheet("Data Murid")
judul(ws, "Data Murid", "Isi sel kuning. Baris 5 adalah contoh — ganti.")
kol = ["No", "Nama", "Kelas asal", "Sekolah", "Kontak orang tua (WA)", "Nama orang tua",
       "Skor diagnostik Matematika (0–16)", "Skor diagnostik BI (0–10)", "Penempatan (otomatis)"]
for j, k in enumerate(kol, 1):
    kepala(ws.cell(4, j), k)
for i in range(N):
    r = R0 + i
    ws.cell(r, 1, i + 1).border = KOTAK
    for j in range(2, 9):
        c = ws.cell(r, j)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.border = KOTAK
    ws.cell(r, 9, f'=IF(G{r}="","",IF(G{r}>=12,"Alur biasa",IF(G{r}>=7,"Alur biasa + soal tambal","Bicarakan dengan orang tua")))').border = KOTAK
contoh = ["Cut Nabila (contoh)", "6B", "SD Negeri Lampineung", "0812-xxxx-xxxx", "Ibu Rahmah", 10, 7]
for j, v in enumerate(contoh, 2):
    ws.cell(R0, j, v).font = Font(name=F, italic=True, color="7F7F7F")
dv = DataValidation(type="whole", operator="between", formula1=0, formula2=16, allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G{R0}:G{R0 + N - 1}")
dv2 = DataValidation(type="whole", operator="between", formula1=0, formula2=10, allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"H{R0}:H{R0 + N - 1}")
ws.cell(R0 + N + 1, 1, "Ambang penempatan Matematika [ASUMSI — belum divalidasi]: 12–16 alur biasa · 7–11 + soal tambal · 0–6 bicarakan dengan orang tua.").font = Font(name=F, size=9, italic=True)
for j, w in enumerate([5, 26, 10, 26, 20, 18, 16, 14, 28], 1):
    ws.column_dimensions[L(j)].width = w
ws.row_dimensions[4].height = 42
ws.freeze_panes = "C5"

# ---------------------------------------------------------------- Tiket Keluar
ws = wb.create_sheet("Tiket Keluar")
judul(ws, "Tiket Keluar", "Skor 0–3 per pertemuan. Kosongkan kalau anak tidak hadir. 3 hijau · 2 kuning · 0–1 merah.")
kepala(ws.cell(3, 1), "Pertemuan")
kepala(ws.cell(4, 1), "Nama murid")
for P in range(1, 37):
    kode, t = JADWAL[P]
    kepala(ws.cell(3, P + 1), f"P{P:02d}")
    kepala(ws.cell(4, P + 1), f"{kode} · {t}", putar=True)
    ws.column_dimensions[L(P + 1)].width = 5.5
ws.row_dimensions[4].height = 150
nama_murid(ws)
rng = f"B{R0}:{L(37)}{R0 + N - 1}"
for i in range(N):
    for P in range(1, 37):
        c = ws.cell(R0 + i, P + 1)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.border = KOTAK
        c.alignment = Alignment(horizontal="center")
dv = DataValidation(type="whole", operator="between", formula1=0, formula2=3, allow_blank=True,
                    error="Isi 0, 1, 2, atau 3", showErrorMessage=True)
ws.add_data_validation(dv)
dv.add(rng)
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(B{R0}<>"",B{R0}=3)'], fill=PatternFill("solid", fgColor=HIJAU)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(B{R0}<>"",B{R0}=2)'], fill=PatternFill("solid", fgColor=KUNING)))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(B{R0}<>"",B{R0}<=1)'], fill=PatternFill("solid", fgColor=MERAH)))
ws.column_dimensions["A"].width = 24
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Penguasaan Topik
ws = wb.create_sheet("Penguasaan Topik")
judul(ws, "Penguasaan Topik",
      "Hijau hanya kalau anak masih bisa mengerjakan topik itu ±3 minggu kemudian (uji tunda). Isi hanya kolom 'Uji tunda'.")
kepala(ws.cell(3, 1), "Topik")
kepala(ws.cell(4, 1), "Nama murid")
STATUS_COL = {}
for t_i, kode in enumerate(TOPIK):
    c0 = 2 + 3 * t_i
    P = P_OF[kode]
    ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 2)
    kepala(ws.cell(3, c0), f"{kode} · {JADWAL[P][1]}")
    for dc in (1, 2):
        ws.cell(3, c0 + dc).border = KOTAK
    kepala(ws.cell(4, c0), f"Tiket P{P:02d}")
    kepala(ws.cell(4, c0 + 1), f"Uji tunda {UJI[kode]}")
    kepala(ws.cell(4, c0 + 2), "Status")
    STATUS_COL[kode] = c0 + 2
    for dc, w in ((0, 7), (1, 10), (2, 10)):
        ws.column_dimensions[L(c0 + dc)].width = w
    tl = L(P + 1)
    for i in range(N):
        r = R0 + i
        a, b, s = L(c0), L(c0 + 1), L(c0 + 2)
        ws.cell(r, c0, f"=IF('Tiket Keluar'!{tl}{r}=\"\",\"\",'Tiket Keluar'!{tl}{r})").border = KOTAK
        u = ws.cell(r, c0 + 1)
        u.fill = PatternFill("solid", fgColor=INPUT)
        u.border = KOTAK
        ws.cell(r, c0 + 2, f'=IF({b}{r}<>"",IF({b}{r}=3,"Kuasai",IF({b}{r}=2,"Baru bisa","Belum")),'
                           f'IF({a}{r}="","",IF({a}{r}>=2,"Baru bisa","Belum")))').border = KOTAK
nama_murid(ws)
last = 2 + 3 * len(TOPIK)
ringkas = [("Jumlah kuasai", "Kuasai", HIJAU), ("Jumlah baru bisa", "Baru bisa", KUNING), ("Jumlah belum", "Belum", MERAH)]
for k, (lab, val, warna) in enumerate(ringkas):
    col = last + k
    kepala(ws.cell(4, col), lab)
    ws.column_dimensions[L(col)].width = 10
    for i in range(N):
        r = R0 + i
        parts = ",".join(f'{L(STATUS_COL[t])}{r}' for t in TOPIK)
        ws.cell(r, col, "=" + "+".join(f'COUNTIF({L(STATUS_COL[t])}{r},"{val}")' for t in TOPIK)).border = KOTAK
area = f"B{R0}:{L(last - 1)}{R0 + N - 1}"
ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"Kuasai"'], fill=PatternFill("solid", fgColor=HIJAU)))
ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"Baru bisa"'], fill=PatternFill("solid", fgColor=KUNING)))
ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"Belum"'], fill=PatternFill("solid", fgColor=MERAH)))
dv = DataValidation(type="whole", operator="between", formula1=0, formula2=3, allow_blank=True)
ws.add_data_validation(dv)
for kode in TOPIK:
    col = L(STATUS_COL[kode] - 1)
    dv.add(f"{col}{R0}:{col}{R0 + N - 1}")
ws.row_dimensions[3].height = 30
ws.row_dimensions[4].height = 45
ws.column_dimensions["A"].width = 24
ws.freeze_panes = "B5"
ws.cell(R0 + N + 1, 1, "Uji tunda M16–M22 dan B08–B11 diambil dari butir simulasi akhir (skala 0–3, lihat 02-peta-spiral.md Bagian 6). "
                        "M22 dan B11 punya dasar paling lemah — sampaikan sebagai 'baru bisa' bila ragu.").font = Font(name=F, size=9, italic=True)

# ---------------------------------------------------------------- Kehadiran
ws = wb.create_sheet("Kehadiran")
judul(ws, "Kehadiran", "H = hadir · T = tidak hadir. Kosong = pertemuan belum berlangsung.")
kepala(ws.cell(4, 1), "Nama murid")
for P in range(1, 37):
    kepala(ws.cell(4, P + 1), f"P{P:02d}")
    ws.column_dimensions[L(P + 1)].width = 4.5
kepala(ws.cell(4, 38), "Hadir")
kepala(ws.cell(4, 39), "Pertemuan tercatat")
kepala(ws.cell(4, 40), "Persentase hadir")
nama_murid(ws)
for i in range(N):
    r = R0 + i
    for P in range(1, 37):
        c = ws.cell(r, P + 1)
        c.fill = PatternFill("solid", fgColor=INPUT)
        c.border = KOTAK
        c.alignment = Alignment(horizontal="center")
    ws.cell(r, 38, f'=COUNTIF(B{r}:AK{r},"H")').border = KOTAK
    ws.cell(r, 39, f'=COUNTIF(B{r}:AK{r},"H")+COUNTIF(B{r}:AK{r},"T")').border = KOTAK
    p = ws.cell(r, 40, f'=IF(AM{r}=0,"",AL{r}/AM{r})')
    p.number_format = "0%"
    p.border = KOTAK
dv = DataValidation(type="list", formula1='"H,T"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"B{R0}:AK{R0 + N - 1}")
ws.conditional_formatting.add(f"B{R0}:AK{R0 + N - 1}", CellIsRule(operator="equal", formula=['"T"'], fill=PatternFill("solid", fgColor=MERAH)))
for c, w in (("A", 24), ("AL", 8), ("AM", 11), ("AN", 11)):
    ws.column_dimensions[c].width = w
ws.row_dimensions[4].height = 32
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Survei Murid
ws = wb.create_sheet("Survei Murid")
judul(ws, "Survei Murid", "Setiap 2 minggu. Pertanyaan 1–3: angka 1–5. Pertanyaan 4: tulis kalimat anak apa adanya. Anonim bila memungkinkan.")
PERT = ["1. Minggu ini aku paham penjelasan gurunya. (1–5)",
        "2. Kalau aku salah, aku tahu letak salahnya di mana. (1–5)",
        "3. Aku berani bertanya kalau tidak mengerti. (1–5)",
        "4. Satu hal yang bikin aku masih bingung:"]
kepala(ws.cell(4, 1), "Nama murid")
ws.column_dimensions["A"].width = 24
putaran = [2, 4, 6, 8, 10, 12]
for k, mg in enumerate(putaran):
    c0 = 2 + 4 * k
    ws.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 3)
    kepala(ws.cell(3, c0), f"Minggu {mg}")
    for q in range(4):
        kepala(ws.cell(4, c0 + q), PERT[q])
        ws.column_dimensions[L(c0 + q)].width = 12 if q < 3 else 30
        if q == 2:
            ws.cell(4, c0 + q).fill = PatternFill("solid", fgColor="F4B084")
    for i in range(N):
        for q in range(4):
            c = ws.cell(R0 + i, c0 + q)
            c.fill = PatternFill("solid", fgColor=INPUT)
            c.border = KOTAK
    rr = R0 + N
    ws.cell(rr, 1, "Rata-rata").font = Font(name=F, bold=True)
    for q in range(3):
        col = L(c0 + q)
        c = ws.cell(rr, c0 + q, f'=IF(COUNT({col}{R0}:{col}{R0 + N - 1})=0,"",AVERAGE({col}{R0}:{col}{R0 + N - 1}))')
        c.number_format = "0.0"
        c.font = Font(name=F, bold=True)
        c.border = KOTAK
    col3 = L(c0 + 2)
    ws.conditional_formatting.add(f"{col3}{rr}", FormulaRule(formula=[f'AND({col3}{rr}<>"",{col3}{rr}<4)'],
                                                            fill=PatternFill("solid", fgColor=MERAH)))
nama_murid(ws)
dv = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True)
ws.add_data_validation(dv)
for k in range(len(putaran)):
    c0 = 2 + 4 * k
    dv.add(f"{L(c0)}{R0}:{L(c0 + 2)}{R0 + N - 1}")
ws.row_dimensions[4].height = 60
cat = R0 + N + 2
ws.cell(cat, 1, "CATATAN: pertanyaan nomor 3 yang paling menentukan. Kelas yang skor nomor 3-nya rendah akan kehilangan murid "
                "dalam sebulan, berapa pun bagusnya skor nomor 1. [sumber: Tahap 2B]").font = Font(name=F, bold=True, size=10)
ws.cell(cat + 1, 1, "Rata-rata nomor 3 di bawah 4 diberi warna merah (petunjuk 05 Bagian 7.4) [ASUMSI — ambang belum divalidasi]. Bahas di rapat mutu minggu itu.").font = Font(name=F, italic=True, size=9)
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Laporan Orang Tua
ws = wb.create_sheet("Laporan Orang Tua")
judul(ws, "Laporan Orang Tua", "Terisi otomatis. Salin kolom 'Pesan WhatsApp' ke orang tua setiap 3 minggu. Tanpa peringkat, tanpa perbandingan.")
kol = ["Nama murid", "Orang tua", "Kehadiran", "Topik yang sudah dikuasai", "Topik yang masih digarap",
       "Saran untuk rumah", "Pesan WhatsApp (siap salin)"]
for j, k in enumerate(kol, 1):
    kepala(ws.cell(4, j), k)
for j, w in enumerate([22, 16, 11, 45, 45, 45, 70], 1):
    ws.column_dimensions[L(j)].width = w
NAMA_SINGKAT = {k: JADWAL[P_OF[k]][1] for k in TOPIK}
pt = "'Penguasaan Topik'"
for i in range(N):
    r = R0 + i
    ws.cell(r, 1, f"=IF('Data Murid'!B{r}=\"\",\"\",'Data Murid'!B{r})")
    ws.cell(r, 2, f"=IF('Data Murid'!F{r}=\"\",\"\",'Data Murid'!F{r})")
    k = ws.cell(r, 3, f"=IF(Kehadiran!AN{r}=\"\",\"\",Kehadiran!AN{r})")
    k.number_format = "0%"
    kuasai = ",".join(f'IF({pt}!{L(STATUS_COL[t])}{r}="Kuasai","{NAMA_SINGKAT[t]}","")' for t in TOPIK)
    garap = ",".join(f'IF(OR({pt}!{L(STATUS_COL[t])}{r}="Baru bisa",{pt}!{L(STATUS_COL[t])}{r}="Belum"),"{NAMA_SINGKAT[t]}","")' for t in TOPIK)
    ws.cell(r, 4, f'=_xlfn.TEXTJOIN(", ",TRUE,{kuasai})')
    ws.cell(r, 5, f'=_xlfn.TEXTJOIN(", ",TRUE,{garap})')
    ws.cell(r, 6, f'=IF(A{r}="","",IF({pt}!{L(last + 2)}{r}>0,'
                  f'"Kalau "&A{r}&" bertanya soal di rumah, tanyakan dulu: kamu sudah coba yang mana? Biarkan "&A{r}&" mencoba sendiri sebelum dibantu.",'
                  f'"Tanyakan ke "&A{r}&" satu hal yang ia pelajari minggu ini, lalu dengarkan tanpa mengoreksi."))')
    ws.cell(r, 7, f'=IF(A{r}="","","Assalamualaikum, "&B{r}&". Laporan 3 mingguan untuk "&A{r}&". '
                  f'Sudah dikuasai: "&IF(D{r}="","belum ada yang diuji ulang",D{r})&". '
                  f'Masih kami kuatkan: "&IF(E{r}="","-",E{r})&". '
                  f'Kehadiran: "&IF(C{r}="","-",TEXT(C{r},"0%"))&". "&F{r})')
    for j in range(1, 8):
        ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, j).border = KOTAK
    ws.row_dimensions[r].height = 90
ws.cell(R0 + N + 1, 1, "Sudah dikuasai = hanya topik hijau (lolos uji tunda). Jangan menambahkan nilai, peringkat, atau jam belajar ke pesan.").font = Font(name=F, italic=True, size=9)
ws.freeze_panes = "B5"

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            f0 = c.font
            c.font = Font(name=F, size=f0.size or 10, bold=f0.bold, italic=f0.italic, color=f0.color)
wb.save(OUT)
print("tersimpan", OUT)
