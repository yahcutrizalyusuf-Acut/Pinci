"""Pengubah markdown sederhana -> python-docx, dipakai bangun.py.

Mendukung: judul #..####, paragraf, **tebal**, *miring*, `kode`, tabel pipa, blok kode ```,
kutipan > (blok kunci), daftar - dan 1., garis ---, baris kosong.
"""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

FONT = "Calibri"
MONO = "Consolas"


def dokumen_baru(ukuran_dasar=12, margin_cm=2.0):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.orientation = WD_ORIENT.PORTRAIT
    for sisi in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, sisi, Cm(margin_cm))
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(ukuran_dasar)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15
    ukuran = {1: ukuran_dasar + 10, 2: ukuran_dasar + 6, 3: ukuran_dasar + 3, 4: ukuran_dasar + 1}
    for lv, sz in ukuran.items():
        h = d.styles[f"Heading {lv}"]
        h.font.name = FONT
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.paragraph_format.space_before = Pt(10 if lv > 1 else 0)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    return d


def nomor_halaman(d, teks_kiri=""):
    for s in d.sections:
        p = s.footer.paragraphs[0]
        p.text = ""
        if teks_kiri:
            r = p.add_run(teks_kiri + "   ·   halaman ")
            r.font.size = Pt(9)
        r = p.add_run()
        for tag, isi in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = isi
            r._r.append(el)
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


PPR_SETELAH_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
                   "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                   "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
                   "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                   "w:sectPr", "w:pPrChange")


def arsir(el_pr, warna):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), warna)
    if el_pr.tag == qn("w:pPr"):
        el_pr.insert_element_before(shd, *PPR_SETELAH_SHD)
    else:
        el_pr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def bingkai_paragraf(p, warna="999999"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for sisi in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{sisi}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), warna)
        bdr.append(e)
    pPr.insert_element_before(bdr, "w:shd", *PPR_SETELAH_SHD)


INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*(?!\s)[^*]+?\*(?![\w*]))")


def tambah_inline(p, teks, ukuran=None, tebal=False, miring=False):
    teks = teks.replace("\\|", "|")
    for bagian in INLINE.split(teks):
        if not bagian:
            continue
        b, m, kode = tebal, miring, False
        if bagian.startswith("***") and bagian.endswith("***") and len(bagian) > 6:
            bagian, b, m = bagian[3:-3], True, True
        elif bagian.startswith("**") and bagian.endswith("**") and len(bagian) > 4:
            bagian, b = bagian[2:-2], True
        elif bagian.startswith("`") and bagian.endswith("`") and len(bagian) > 2:
            bagian, kode = bagian[1:-1], True
        elif bagian.startswith("*") and bagian.endswith("*") and len(bagian) > 2:
            bagian, m = bagian[1:-1], True
        r = p.add_run(bagian)
        r.bold, r.italic = b, m
        if kode:
            r.font.name = MONO
        if ukuran:
            r.font.size = Pt(ukuran)
    return p


def sel_tabel(baris):
    baris = baris.strip()
    if baris.startswith("|"):
        baris = baris[1:]
    if baris.endswith("|"):
        baris = baris[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", baris)]


def tabel(d, baris_tabel, ukuran=10.5, lebar_cm=17.0, kepala=True):
    rows = [sel_tabel(b) for b in baris_tabel if not re.fullmatch(r"\|?[\s:\-|]+\|?", b.strip())]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = d.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # lebar kolom menurut panjang isi
    panjang = [max(3, min(40, max(len(r[i]) if i < len(r) else 0 for r in rows))) for i in range(ncol)]
    tot = sum(panjang)
    lebar = [Cm(lebar_cm * p / tot) for p in panjang]
    for i, r in enumerate(rows):
        tr = t.rows[i]
        trPr = tr._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for j in range(ncol):
            c = tr.cells[j]
            c.width = lebar[j]
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            isi = r[j] if j < len(r) else ""
            tambah_inline(c.paragraphs[0], isi, ukuran=ukuran, tebal=(kepala and i == 0))
            if kepala and i == 0:
                arsir(c._tc.get_or_add_tcPr(), "D9D9D9")
    if kepala:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        trPr.append(th)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def blok_kode(d, baris, ukuran=10.5):
    for b in baris:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(b if b else " ")
        r.font.name = MONO
        r.font.size = Pt(ukuran)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def tulis_markdown(d, teks, ukuran=12, pecah_halaman_pada=None, lewati_judul1=False, ukuran_kunci=10.5,
                   geser_judul=0):
    """pecah_halaman_pada: regex judul '## ...' yang diawali halaman baru."""
    baris = teks.splitlines()
    i = 0
    pertama_h2 = True
    while i < len(baris):
        b = baris[i]
        s = b.strip()
        if s.startswith("```"):
            j = i + 1
            isi = []
            while j < len(baris) and not baris[j].strip().startswith("```"):
                isi.append(baris[j])
                j += 1
            blok_kode(d, isi)
            i = j + 1
            continue
        if s.startswith("|"):
            j = i
            tb = []
            while j < len(baris) and baris[j].strip().startswith("|"):
                tb.append(baris[j])
                j += 1
            tabel(d, tb, ukuran=ukuran - 1.5)
            i = j
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lv = min(4, len(m.group(1)) + geser_judul)
            judul = m.group(2)
            if lv == 1 and lewati_judul1:
                i += 1
                continue
            if pecah_halaman_pada and len(m.group(1)) == 2 and re.match(pecah_halaman_pada, judul):
                if not pertama_h2:
                    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                pertama_h2 = False
            h = d.add_heading(level=lv)
            tambah_inline(h, judul)
            i += 1
            continue
        if s.startswith(">"):
            j = i
            while j < len(baris) and baris[j].strip().startswith(">"):
                isi = baris[j].strip()[1:].strip()
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                pPr = p._p.get_or_add_pPr()
                arsir(pPr, "F2F2F2")
                tambah_inline(p, isi, ukuran=ukuran_kunci)
                j += 1
            d.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j
            continue
        if s in ("---", "***", "==="):
            i += 1
            continue
        if not s:
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", b)
        if m and not re.match(r"^[A-E]\.\s", s):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (len(m.group(1)) // 2))
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(2)
            tanda = "–  " if m.group(2) in "-*" else m.group(2) + "  "
            p.add_run(tanda)
            tambah_inline(p, m.group(3), ukuran=ukuran)
            i += 1
            continue
        p = d.add_paragraph()
        tambah_inline(p, s, ukuran=ukuran)
        i += 1
